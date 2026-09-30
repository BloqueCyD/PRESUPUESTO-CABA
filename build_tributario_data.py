#!/usr/bin/env python3
"""
build_tributario_data.py — genera data/tributario/*.json para la solapa "Tributario".

Uso:
    python3 build_tributario_data.py

Sólo usa la biblioteca estándar de Python (3.9+) más openpyxl (ya usado para
leer/escribir la planilla de referencia del equipo).

Fuentes (en fuentes_tributario/, sí versionadas: son chicas):
  * ley_impositiva_caba_base_2026.csv — una fila por (año, concepto), con
    artículo, fuente oficial y estado "a verificar" de cada valor. Es el
    registro de trabajo del equipo, cargado a mano con lectura de los PDF
    oficiales de la Ley Impositiva / Tarifaria y el Código Fiscal 2023-2026.
  * plantilla_ley_impositiva_caba.xlsx, hoja "Plantilla" — el mismo universo
    de conceptos, pero ya agrupado a través de los años (una fila por
    concepto, con su valor en cada año 2023-2026), resolviendo los cambios de
    redacción/numeración de artículo que usan los mismos conceptos. Es la
    fuente de verdad para "qué valor de 2023 corresponde al mismo concepto
    que este valor de 2024/2025/2026".

Esta build NO vuelve a resolver esa agrupación: la toma de la Plantilla y
re-vincula cada celda con su artículo/fuente original del CSV largo por
coincidencia de texto (impuesto + año + tipo de valor + concepto,
comparación insensible a mayúsculas/acentos/orden de palabras y a
paréntesis aclaratorios). Una celda que no encuentra una coincidencia
confiable (score < UMBRAL) queda con fuente=null: se muestra el valor igual,
pero sin cita de artículo — nunca se inventa una cita.

Salida: data/tributario/package.json y data/tributario/tax-values.json,
consumidos por app.js en la solapa Tributario. Alícuotas y montos, 2023-2026
(no incluye Código Fiscal, Ley Arancelaria ni Agenda prioritaria: esas vistas
quedan como "próximamente" hasta que haya una extracción propia).

IPC promedio anual: las mismas constantes (fuente IPCBA / pauta Presupuesto
Nacional) que usa build_data.py para el módulo de Presupuesto, así que la
variación real es comparable entre ambos módulos.
"""
import csv
import difflib
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")

HERE = Path(__file__).resolve().parent
FUENTES = HERE / "fuentes_tributario"
CSV_PATH = FUENTES / "ley_impositiva_caba_base_2026.csv"
XLSX_PATH = FUENTES / "plantilla_ley_impositiva_caba.xlsx"
OUT_DIR = HERE / "data" / "tributario"

YEARS = ["2023", "2024", "2025", "2026"]
COLMAP = {"2023": "Valor 2023", "2024": "Valor 2024", "2025": "Valor 2025", "2026": "Valor 2026"}
MATCH_THRESHOLD = 0.78

# Mismas constantes que build_data.py (fuente: IPCBA, ver DICCIONARIO.md).
IPC_PROMEDIO = {"2023": 397.16, "2024": 1295.33, "2025": 1896.37, "2026": 2501.96}

TAX_ORDER = [
    "Ingresos Brutos", "Inmobiliario", "Patentes", "Sellos", "Publicidad",
    "Antenas", "Espacio Público", "Electricidad", "Régimen Simplificado", "Ambientales",
]
VALUE_TYPE_LABEL = {
    "alicuota_general": "Alícuota general", "alicuota_diferencial": "Alícuota diferencial",
    "monto_fijo": "Monto fijo", "minimo": "Mínimo", "tope": "Tope", "recargo": "Recargo",
    "umbral_exencion": "Umbral de exención", "umbral_condonacion": "Umbral de condonación",
    "mecanismo": "Mecanismo de cálculo", "informativo": "Informativo",
}


def strip_accents(s):
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def normalize(s):
    if not s:
        return ""
    s = strip_accents(s).lower()
    s = re.sub(r"[^a-z0-9%\s/.\-]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def normalize_for_match(s):
    return normalize(re.sub(r"\([^)]*\)", " ", s))


STOP = {"", "de", "la", "el", "y", "a", "en", "por"}


def tokens(s):
    return set(re.split(r"[\s/\-]+", s)) - STOP


def match_score(a, b):
    na, nb = normalize_for_match(a), normalize_for_match(b)
    if not na or not nb:
        return 0.0
    if na == nb:
        return 1.0
    r = difflib.SequenceMatcher(None, na, nb).ratio()
    if na in nb or nb in na:
        r = max(r, 0.85)
    ta, tb = tokens(na), tokens(nb)
    if ta and tb:
        jac = len(ta & tb) / len(ta | tb)
        sr = difflib.SequenceMatcher(None, " ".join(sorted(ta)), " ".join(sorted(tb))).ratio()
        r = max(r, jac, sr * 0.95 if jac >= 0.5 else 0)
    return r


def load_long_csv():
    with open(CSV_PATH, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    by_key = {}
    for r in rows:
        by_key.setdefault((r["impuesto"], r["anio"], r["tipo_valor"]), []).append(r)
    return by_key


def load_plantilla():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    ws = wb["Plantilla"]
    headers = [c.value for c in ws[2]]
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(headers, row)))
    return rows


def build_concepts(plantilla_rows, by_key):
    concepts = []
    seen_ids = {}
    matched = total = 0
    for row in plantilla_rows:
        impuesto, concepto = row["Impuesto"], row["Concepto"]
        tipo, unidad = row["Tipo de valor"], row["Unidad"]
        nota = (row["Nota / a verificar"] or "").strip()

        cid_base = re.sub(r"-+", "-", normalize(f"{impuesto}-{concepto}-{tipo}").replace(" ", "-"))[:80]
        n = seen_ids.get(cid_base, 0)
        seen_ids[cid_base] = n + 1
        cid = cid_base if n == 0 else f"{cid_base}-{n}"

        values, refs = {}, {}
        pending = "[verificar]" in nota.lower()
        for y in YEARS:
            v = row[COLMAP[y]]
            if v is None or v == "":
                values[y] = None
                refs[y] = None
                continue
            values[y] = v
            total += 1
            best, best_score = None, 0.0
            for c in by_key.get((impuesto, y, tipo), []):
                sc = match_score(concepto, c["concepto"])
                if sc > best_score:
                    best, best_score = c, sc
            if best and best_score >= MATCH_THRESHOLD:
                matched += 1
                refs[y] = {
                    "articulo": best["articulo"] or None,
                    "fuente": best["fuente"] or None,
                }
                if best["verificar"] == "SI":
                    pending = True
            else:
                refs[y] = None

        concepts.append({
            "conceptId": cid,
            "tax": impuesto,
            "concept": concepto,
            "valueType": tipo,
            "valueTypeLabel": VALUE_TYPE_LABEL.get(tipo, tipo),
            "unit": unidad,
            "values": values,
            "refs": refs,
            "note": nota or None,
            "reviewStatus": "pending_review" if pending else "confirmed",
        })
    return concepts, matched, total


def main():
    if not CSV_PATH.exists() or not XLSX_PATH.exists():
        sys.exit(f"Faltan fuentes en {FUENTES}")

    by_key = load_long_csv()
    plantilla_rows = load_plantilla()
    concepts, matched, total = build_concepts(plantilla_rows, by_key)

    taxes_present = sorted({c["tax"] for c in concepts}, key=lambda t: TAX_ORDER.index(t) if t in TAX_ORDER else 99)
    pending_count = sum(1 for c in concepts if c["reviewStatus"] == "pending_review")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "tax-values.json").write_text(
        json.dumps({"concepts": concepts}, ensure_ascii=False, indent=None, separators=(",", ":")),
        encoding="utf-8",
    )

    package = {
        "packageId": "caba-tributario-2023-2026",
        "generatedAt": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "status": "preliminar",
        "years": [int(y) for y in YEARS],
        "defaultBaseYear": 2025,
        "defaultComparisonYear": 2026,
        "proyecto2027Disponible": False,
        "ipcPromedio": {int(y): v for y, v in IPC_PROMEDIO.items()},
        "taxes": taxes_present,
        "documents": [
            {"id": f"impositiva-{y}", "type": "impositiva", "year": int(y), "status": "processed"}
            for y in YEARS
        ],
        "coverage": {
            "conceptos": len(concepts),
            "valoresTotales": total,
            "valoresConFuenteVinculada": matched,
            "conceptosPendientesDeRevision": pending_count,
        },
        "views": {
            "resumen": "proximamente",
            "impositiva": "disponible",
            "codigo_fiscal": "proximamente",
            "arancelaria": "proximamente",
            "agenda": "proximamente",
        },
    }
    (OUT_DIR / "package.json").write_text(
        json.dumps(package, ensure_ascii=False, indent=2), encoding="utf-8",
    )

    print(f"OK: {len(concepts)} conceptos, {total} valores-año, "
          f"{matched} con fuente vinculada ({matched / total * 100:.1f}%), "
          f"{pending_count} conceptos con algo a verificar.")
    print(f"Escrito: {OUT_DIR / 'tax-values.json'}")
    print(f"Escrito: {OUT_DIR / 'package.json'}")


if __name__ == "__main__":
    main()
